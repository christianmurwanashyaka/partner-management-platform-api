import os
from datetime import datetime
from io import BytesIO
import platform
import subprocess
from tempfile import NamedTemporaryFile

from docx2pdf import convert
from fastapi import status

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt
from fastapi import UploadFile, HTTPException, Depends
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from sqlmodel import select
import pypandoc


from db.database import get_db
from db.models import Activity, ActivityDomain, InputDetail, MouApplication, MouDetail, Project, OrganizationType, \
    Party, Goal, FundingSource, FundingUnit, Organization, BudgetType, Report, ReportActivity


async def handle_upload_file(file: UploadFile):
    upload_directory = os.path.join(os.getcwd(), 'uploads')
    os.makedirs(upload_directory, exist_ok=True)

    try:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f'{timestamp}_{file.filename}'
        file_path = os.path.join(upload_directory, filename)

        with open(file_path, 'wb+') as file_object:
            file_object.write(await file.read())

        return file_path, filename
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to save file: {str(e)}")


async def generate_mou_doc(mou_application, template_path, db: AsyncSession = Depends(get_db)):
    doc = Document(template_path)

    mou_detail_query = select(MouDetail).where(MouDetail.uuid == mou_application.mou_detail_id)
    mou_detail = (await db.execute(mou_detail_query)).scalar_one_or_none()

    parties_query = select(Party).where(Party.mou_detail_id == mou_detail.uuid)
    parties = (await db.execute(parties_query)).scalars().all()

    project_query = select(Project).where(Project.uuid == mou_detail.project_id).options(selectinload(Project.organization))
    project = (await db.execute(project_query)).scalar_one_or_none()

    activities_query = select(Activity).where(Activity.project_id == project.uuid).options(selectinload(Activity.domains))
    activities = (await db.execute(activities_query)).scalars().all()

    goals_query = select(Goal).where(Goal.project_id == project.uuid)
    goals = (await db.execute(goals_query)).scalars().all()

    organization = project.organization
    current_date = datetime.now().strftime("%d/%m/%Y")

    def get_attr_or_none(obj, attr):
        value = getattr(obj, attr, None)
        return value if value != '' else None

    home_country = get_attr_or_none(organization, 'home_country')
    home_country_province = get_attr_or_none(organization, 'home_country_province')
    home_country_district = get_attr_or_none(organization, 'home_country_district')
    home_country_avenue = get_attr_or_none(organization, 'home_country_avenue')
    home_country_po_box = get_attr_or_none(organization, 'home_country_po_box')

    # Collect unique domains from the project's activities
    domains = set()
    for activity in project.activities:
        for activity_domain in activity.domains:
            if activity_domain.domain_intervention:
                domains.add(activity_domain.domain_intervention.name)
    domains_str = ', '.join(domains)

    # Collect responsibilities and signatory details for the party with organization_id
    responsibilities_str = ""
    party_signatory = ""
    party_position = "CEO"  # Default value
    for party in parties:
        if party.organization_id:
            responsibilities_str = "\n".join(
                [f"{idx + 1}. {responsibility}" for idx, responsibility in enumerate(party.responsibilities)])
            party_signatory = party.signatory
            party_position = party.position

    project_duration = project.duration if project.duration else "1"

    mappings = {
        '{ORGANIZATION_NAME}': organization.name,
        '{HOME_COUNTRY}': home_country,
        '{HOME_COUNTRY_PROVINCE}': home_country_province,
        '{HOME_COUNTRY_DISTRICT}': home_country_district,
        '{HOME_COUNTRY_AVENUE}': home_country_avenue,
        '{HOME_COUNTRY_PO_BOX}': home_country_po_box,
        '{PROJECT_NAME}': project.name,
        '{PROJECT_DURATION}': project_duration,
        '{ORGANIZATION_PO_BOX}': organization.rwanda_po_box,
        '{ORGANIZATION_PHONE}': organization.phone_number,
        '{ORGANIZATION_EMAIL}': organization.email,
        '{ORGANIZATION_WEBSITE}': organization.website,
        '{ORGANIZATION_ADDRESS}': organization.rwanda_avenue,
        '{OVERALL_GOAL}': project.overall_goal,
        '{ACTIVITIES_DOMAINS}': domains_str,
        '{PARTY_RESPONSIBILITIES}': responsibilities_str,
        '{PARTY_SIGNATORY_NAME}': party_signatory,
        '{PARTY_SIGNATORY_POSITION}': party_position,
        '{DATE}': current_date,
    }

    first_occurrence = True

    def replace_text(element, mappings):
        nonlocal first_occurrence
        for key, value in mappings.items():
            if key in element.text:
                if key == '{ORGANIZATION_NAME}' and first_occurrence:
                    # Split the text into parts
                    parts = element.text.split(key)
                    element.text = parts[0]

                    # Add the organization name with uppercase and bold formatting
                    run = element.add_run(value.upper())
                    run.bold = True

                    # Add any remaining text
                    element.add_run(parts[1] if len(parts) > 1 else '')

                    first_occurrence = False
                else:
                    element.text = element.text.replace(key, str(value))

    # Iterate through all elements to replace placeholders
    for paragraph in doc.paragraphs:
        replace_text(paragraph, mappings)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_text(paragraph, mappings)

    for section in doc.sections:
        header = section.header
        footer = section.footer
        for paragraph in header.paragraphs:
            replace_text(paragraph, mappings)
        for paragraph in footer.paragraphs:
            replace_text(paragraph, mappings)

    if 'List Number' not in [s.name for s in doc.styles]:
        style = doc.styles.add_style('List Number', WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles['Normal']
        style.paragraph_format.left_indent = Pt(36)  # Set the left indentation
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.font.name = 'Times New Roman'
        style.font.size = doc.styles['Normal'].font.size

        p = style.element
        num = OxmlElement('w:numPr')
        ilvl = OxmlElement('w:ilvl')
        ilvl.set(qn('w:val'), "0")
        numId = OxmlElement('w:numId')
        numId.set(qn('w:val'), "1")
        num.append(ilvl)
        num.append(numId)
        p.append(num)

    def set_font(paragraph, font_name='Times New Roman'):
        for run in paragraph.runs:
            run.font.name = font_name
            run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

    for paragraph in doc.paragraphs:
        if '{PROJECT_GOALS}' in paragraph.text:
            paragraph.text = paragraph.text.replace('{PROJECT_GOALS}', '')
            for goal in goals:
                goal_paragraph = paragraph.insert_paragraph_before(goal.name)
                goal_paragraph.style = doc.styles['List Number']
                set_font(goal_paragraph)

    docx_buffer = BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)

    pdf_buffer = convert_docx_to_pdf(docx_buffer)

    return docx_buffer, pdf_buffer


def convert_docx_to_pdf(docx_buffer: BytesIO) -> BytesIO:
    pdf_buffer = BytesIO()

    with NamedTemporaryFile(delete=False, suffix='.docx') as tmp_docx:
        tmp_docx.write(docx_buffer.getvalue())
        tmp_docx_path = tmp_docx.name

    try:
        if platform.system() == 'Linux':
            pdf_path = convert_docx_to_pdf_linux(tmp_docx_path)
            if pdf_path:
                with open(pdf_path, 'rb') as pdf_file:
                    pdf_buffer.write(pdf_file.read())
                os.unlink(pdf_path)  # Remove the temporary PDF file
            else:
                raise Exception("PDF conversion failed on Linux")
        else:
            pdf_path = convert_docx_to_pdf_windows_mac(tmp_docx_path)
            with open(pdf_path, 'rb') as pdf_file:
                pdf_buffer.write(pdf_file.read())
            os.unlink(pdf_path)  # Remove the temporary PDF file

    finally:
        os.unlink(tmp_docx_path)  # Always remove the temporary DOCX file

    pdf_buffer.seek(0)
    return pdf_buffer

def convert_docx_to_pdf_linux(docx_file_path: str) -> str:
    pdf_file_path = os.path.splitext(docx_file_path)[0] + '.pdf'
    try:
        subprocess.run(
            ['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir',
             os.path.dirname(pdf_file_path), docx_file_path],
            check=True, capture_output=True, text=True
        )
        return pdf_file_path
    except subprocess.CalledProcessError as e:
        print(f"LibreOffice conversion failed! Error: {e}")
        print(f"STDOUT: {e.stdout}")
        print(f"STDERR: {e.stderr}")
        return None

def convert_docx_to_pdf_windows_mac(docx_file_path: str) -> str:
    from docx2pdf import convert
    pdf_file_path = os.path.splitext(docx_file_path)[0] + '.pdf'
    convert(docx_file_path, pdf_file_path)
    return pdf_file_path

async def generate_mou_action_plan(mou_application, db: AsyncSession = Depends(get_db)):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Action Plan'

    headers = [
        'Organization',
        'Organization Type',
        'Project Name',
        'Domain of Intervention',
        'Sub Domain of Intervention',
        'Sub Domain Function',
        'Sub Function',
        'Location',
        'Funding Source',
        'Funding Unit',
        'Activity',
        'Description of activity',
        'Input Category',
        'Inputs',
        'Planned budget',
        'Currency',
        'On/Off Budget/IGR',
        'Implementer',
        'Fiscal Year'
    ]
    ws.append(headers)

    bold_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold_font
    query = select(MouApplication).options(
        selectinload(MouApplication.mou_detail)
    ).where(MouApplication.uuid == mou_application.uuid)

    result = await db.execute(query)
    mou_application = result.scalar_one_or_none()

    mou_detail_query = select(MouDetail).where(MouDetail.uuid == mou_application.mou_detail_id)
    mou_detail = (await db.execute(mou_detail_query)).scalar_one_or_none()

    project_query = select(Project).options(
        selectinload(Project.funding_source),
        selectinload(Project.funding_unit),
        selectinload(Project.organization),
        selectinload(Project.budget_type),
    ).where(Project.uuid == mou_detail.project_id)
    project = (await db.execute(project_query)).scalar_one_or_none()

    if not mou_application:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail='MouApplication not found')

    funding_source_query = select(FundingSource).where(FundingSource.uuid == project.funding_source_id)
    funding_source_result = (await db.execute(funding_source_query)).scalar_one_or_none()
    funding_source = funding_source_result if funding_source_result else (project.other_funding_source if project.other_funding_source else "N/A")

    budget_type_query = select(BudgetType).where(BudgetType.uuid == project.budget_type_id)
    budget_type = (await db.execute(budget_type_query)).scalar_one_or_none()

    funding_unit_query = select(FundingUnit).where(FundingUnit.uuid == project.funding_unit_id)
    funding_unit_result = (await db.execute(funding_unit_query)).scalar_one_or_none()
    funding_unit = funding_unit_result if funding_unit_result else 'N/A'

    organization_query = select(Organization).where(Organization.uuid == project.organization_id)
    organization = (await db.execute(organization_query)).scalar_one_or_none()

    organization_type_query = select(OrganizationType).where(OrganizationType.uuid == organization.organization_type_id)
    organization_type = (await db.execute(organization_type_query)).scalar_one_or_none()


    # Fetching activities and their related data
    project_activities_query = select(Activity).where(Activity.project_id == project.uuid)
    activities = (await db.execute(project_activities_query)).scalars().all()

    for activity in activities:
        activity_domains_query = select(ActivityDomain).where(ActivityDomain.activity_id == activity.uuid).options(
            selectinload(ActivityDomain.domain_intervention),
            selectinload(ActivityDomain.sub_domain),
            selectinload(ActivityDomain.sub_domain_function),
            selectinload(ActivityDomain.sub_function),
        )
        activity_domains = (await db.execute(activity_domains_query)).scalars().all()

        input_details_query = select(InputDetail).where(InputDetail.activity_id == activity.uuid).options(
            selectinload(InputDetail.input_category),
            selectinload(InputDetail.input),
        )
        input_details = (await db.execute(input_details_query)).scalars().all()

        for input_detail in input_details:
            district = getattr(input_detail, 'district', 'N/A')
            province = getattr(input_detail, 'province', 'N/A')
            location = f"{district}, {province}"
            input_category = getattr(getattr(input_detail, 'input_category', None), 'name', 'N/A')
            input_name = getattr(getattr(input_detail, 'input', None), 'name', 'N/A')
            budget = getattr(input_detail, 'budget', 'N/A')

            for domain in activity_domains:
                domain_name = getattr(getattr(domain, 'domain_intervention', None), 'name', 'N/A')
                sub_domain_name = getattr(getattr(domain, 'sub_domain', None), 'name', 'N/A')
                sub_domain_function_name = getattr(getattr(domain, 'sub_domain_function', None), 'name', 'N/A')
                sub_function_name = getattr(getattr(domain, 'sub_function', None), 'name', 'N/A')

                data = [
                    organization.name,
                    organization_type.name,
                    project.name,
                    domain_name,
                    sub_domain_name,
                    sub_domain_function_name,
                    sub_function_name,
                    location,
                    funding_source.name,
                    funding_unit.name,
                    activity.name,
                    activity.description,
                    input_category,
                    input_name,
                    budget,
                    project.currency,
                    budget_type.name,
                    activity.implementer,
                    activity.fiscal_year
                ]
                ws.append(data)

    action_plans_directory = os.path.join(os.getcwd(), 'action_plans')
    os.makedirs(action_plans_directory, exist_ok=True)
    try:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f'{timestamp}_action_plan_{mou_application.created_by}.xlsx'
        file_path = os.path.join(action_plans_directory, filename)
        wb.save(file_path)
        return file_path, filename
    except Exception as e:
        raise Exception(f"Failed to save file: {str(e)}")


async def save_mou_doc_to_disk(document_buffer, filename, doc_type):
    directory = os.path.join(os.getcwd(), f'generated_mou_docs_{doc_type}')
    os.makedirs(directory, exist_ok=True)
    file_path = os.path.join(directory, filename)
    with open(file_path, 'wb') as file:
        file.write(document_buffer.getvalue())  # Use getvalue() instead of read()
    return file_path, filename


async def generate_implementation_plan(report: Report, db: AsyncSession):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Implementation Plan'

    headers = [
        'Organization',
        'Organization Type',
        'Project Name',
        'Domain of Intervention',
        'Sub Domain of Intervention',
        'Sub Domain Function',
        'Sub Function',
        'Location',
        'Funding Source',
        'Funding Unit',
        'Activity',
        'Description of activity',
        'Input Category',
        'Inputs',
        'Planned budget',
        'Currency',
        'On/Off Budget/IGR',
        'Implementer',
        'Fiscal Year',
        'Executed Budget',
        'Accomplishments'
    ]
    ws.append(headers)

    bold_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold_font

    # Fetch project data
    project_query = select(Project).where(Project.uuid == report.project_uuid)
    project_result = await db.execute(project_query)
    project = project_result.scalar_one_or_none()

    if not project:
        raise HTTPException(status_code=404, detail='Project not found')

    # Fetch organization data
    organization_query = select(Organization).where(Organization.uuid == project.organization_id)
    organization_result = await db.execute(organization_query)
    organization = organization_result.scalar_one_or_none()

    if not organization:
        raise HTTPException(status_code=404, detail='Organization not found')

    # Fetch organization type
    organization_type_query = select(OrganizationType).where(OrganizationType.uuid == organization.organization_type_id)
    organization_type_result = await db.execute(organization_type_query)
    organization_type = organization_type_result.scalar_one_or_none()

    # Fetch funding source
    funding_source_query = select(FundingSource).where(FundingSource.uuid == project.funding_source_id)
    funding_source_result = await db.execute(funding_source_query)
    funding_source = funding_source_result.scalar_one_or_none()
    funding_source_name = funding_source.name if funding_source else project.other_funding_source or "N/A"

    # Fetch funding unit
    funding_unit_query = select(FundingUnit).where(FundingUnit.uuid == project.funding_unit_id)
    funding_unit_result = await db.execute(funding_unit_query)
    funding_unit = funding_unit_result.scalar_one_or_none()
    funding_unit_name = funding_unit.name if funding_unit else 'N/A'

    # Fetch budget type
    budget_type_query = select(BudgetType).where(BudgetType.uuid == project.budget_type_id)
    budget_type_result = await db.execute(budget_type_query)
    budget_type = budget_type_result.scalar_one_or_none()

    # Fetch activities associated with the report
    activities_query = select(Activity).where(Activity.report_uuid == report.uuid)
    activities_result = await db.execute(activities_query)
    activities = activities_result.scalars().all()

    for activity in activities:
        activity_domains_query = select(ActivityDomain).where(ActivityDomain.activity_id == activity.uuid).options(
            selectinload(ActivityDomain.domain_intervention),
            selectinload(ActivityDomain.sub_domain),
            selectinload(ActivityDomain.sub_domain_function),
            selectinload(ActivityDomain.sub_function),
        )
        activity_domains = (await db.execute(activity_domains_query)).scalars().all()

        input_details_query = select(InputDetail).where(InputDetail.activity_id == activity.uuid).options(
            selectinload(InputDetail.input_category),
            selectinload(InputDetail.input),
        )
        input_details = (await db.execute(input_details_query)).scalars().all()

        # Fetch corresponding ReportActivity
        report_activity_query = select(ReportActivity).where(ReportActivity.activity_uuid == activity.uuid)
        report_activity_result = await db.execute(report_activity_query)
        # TODO: UPDATE TO SCALAR_ONE_OR_NONE()
        report_activity = report_activity_result.scalars().all()[0]

        for input_detail in input_details:
            district = getattr(input_detail, 'district', 'N/A')
            province = getattr(input_detail, 'province', 'N/A')
            location = f"{district}, {province}"
            input_category = getattr(getattr(input_detail, 'input_category', None), 'name', 'N/A')
            input_name = getattr(getattr(input_detail, 'input', None), 'name', 'N/A')
            budget = getattr(input_detail, 'budget', 'N/A')

            for domain in activity_domains:
                domain_name = getattr(getattr(domain, 'domain_intervention', None), 'name', 'N/A')
                sub_domain_name = getattr(getattr(domain, 'sub_domain', None), 'name', 'N/A')
                sub_domain_function_name = getattr(getattr(domain, 'sub_domain_function', None), 'name', 'N/A')
                sub_function_name = getattr(getattr(domain, 'sub_function', None), 'name', 'N/A')

                data = [
                    organization.name,
                    organization_type.name,
                    project.name,
                    domain_name,
                    sub_domain_name,
                    sub_domain_function_name,
                    sub_function_name,
                    location,
                    funding_source_name,
                    funding_unit_name,
                    activity.name,
                    activity.description,
                    input_category,
                    input_name,
                    budget,
                    project.currency,
                    budget_type.name if budget_type else 'N/A',
                    activity.implementer,
                    activity.fiscal_year,
                    report_activity.executed_budget if report_activity else 'N/A',
                    ', '.join(report_activity.accomplishments) if report_activity and report_activity.accomplishments else 'N/A'
                ]
                ws.append(data)

    implementation_plans_directory = os.path.join(os.getcwd(), 'implementation_plans')
    os.makedirs(implementation_plans_directory, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f'{timestamp}_implementation_plan_{report.reported_by}.xlsx'
    file_path = os.path.join(implementation_plans_directory, filename)
    wb.save(file_path)

    return file_path, filename